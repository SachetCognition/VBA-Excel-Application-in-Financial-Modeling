"""
Generate Excel workbooks with charts matching all VBA macro outputs.

Creates:
1. LoanAmortization.xlsx — All 3 test cases with 14 charts each
2. PortfolioAnalysis.xlsx — Efficient frontier, Sharpe, VCV, weights
"""

import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import (
    ScatterChart, BarChart, LineChart, Reference
)
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers

from before.loan_amortization_vba_recreation import (
    constant_payment_schedule, straight_line_schedule, summary_statistics
)
from before.portfolio_analysis_vba_recreation import full_portfolio_analysis

TEST_CASES = [
    {"name": "Case1", "principal": 100000, "rate": 5, "years": 30,
     "freq": "Monthly", "ptype": "End of Period",
     "label": "$100K / 5% / 30yr / Monthly / End"},
    {"name": "Case2", "principal": 250000, "rate": 3.5, "years": 15,
     "freq": "Monthly", "ptype": "Begin of Period",
     "label": "$250K / 3.5% / 15yr / Monthly / Begin"},
    {"name": "Case3", "principal": 50000, "rate": 7, "years": 5,
     "freq": "Annually", "ptype": "End of Period",
     "label": "$50K / 7% / 5yr / Annual / End"},
]

HEADER_FILL = PatternFill(start_color="0F3460", end_color="0F3460", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11, name="Calibri")
DATA_FONT = Font(size=10, name="Calibri")
TITLE_FONT = Font(bold=True, size=14, name="Times New Roman")
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)


def write_schedule_to_sheet(ws, df, start_row, title, params_text):
    """Write an amortization schedule to a worksheet with formatting."""
    # Title
    ws.cell(row=start_row, column=1, value=title).font = TITLE_FONT
    ws.merge_cells(start_row=start_row, start_column=1,
                   end_row=start_row, end_column=8)
    ws.cell(row=start_row + 1, column=1, value=params_text).font = Font(
        italic=True, size=10, color="666666")

    # Headers
    headers = ["Period", "BegBal", "Payment", "Interest", "Principal",
               "EndBal", "Interest%", "Principal%"]
    header_row = start_row + 3
    for j, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=j, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center')
        cell.border = THIN_BORDER

    # Data
    for i, row in df.iterrows():
        r = header_row + 1 + i
        for j, col in enumerate(df.columns, 1):
            cell = ws.cell(row=r, column=j, value=row[col])
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            if col in ['BegBal', 'Payment', 'Interest', 'Principal', 'EndBal']:
                cell.number_format = '#,##0.00'
            elif col in ['InterestPct', 'PrincipalPct']:
                cell.number_format = '0.00%'

    # Column widths
    widths = [8, 14, 12, 12, 12, 14, 12, 12]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    return header_row, header_row + len(df)


def add_7_charts(ws, header_row, last_row, title_prefix, chart_col=10):
    """Add 7 VBA-matching charts to the worksheet."""
    n_rows = last_row - header_row
    chart_row = header_row

    # Chart 1: BegBal scatter
    c1 = ScatterChart()
    c1.title = f"{title_prefix}: Beginning Balance"
    c1.x_axis.title = "Periods"
    c1.y_axis.title = "Balance ($)"
    c1.style = 13
    c1.width = 18
    c1.height = 10
    xvals = Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_row)
    yvals = Reference(ws, min_col=2, min_row=header_row + 1, max_row=last_row)
    c1.add_data(yvals)
    c1.set_categories(xvals)
    c1.series[0].title = SeriesLabel(v="BegBal")
    ws.add_chart(c1, f"{get_column_letter(chart_col)}{chart_row}")

    # Chart 2: Payment scatter
    c2 = ScatterChart()
    c2.title = f"{title_prefix}: Payment"
    c2.x_axis.title = "Periods"
    c2.y_axis.title = "Payment ($)"
    c2.style = 13
    c2.width = 18
    c2.height = 10
    yvals2 = Reference(ws, min_col=3, min_row=header_row + 1, max_row=last_row)
    c2.add_data(yvals2)
    c2.set_categories(xvals)
    c2.series[0].title = SeriesLabel(v="Payment")
    ws.add_chart(c2, f"{get_column_letter(chart_col)}{chart_row + 16}")

    # Chart 3: Interest scatter
    c3 = ScatterChart()
    c3.title = f"{title_prefix}: Interest"
    c3.x_axis.title = "Periods"
    c3.y_axis.title = "Interest ($)"
    c3.style = 13
    c3.width = 18
    c3.height = 10
    yvals3 = Reference(ws, min_col=4, min_row=header_row + 1, max_row=last_row)
    c3.add_data(yvals3)
    c3.set_categories(xvals)
    c3.series[0].title = SeriesLabel(v="Interest")
    ws.add_chart(c3, f"{get_column_letter(chart_col)}{chart_row + 32}")

    # Chart 4: Principal scatter
    c4 = ScatterChart()
    c4.title = f"{title_prefix}: Principal"
    c4.x_axis.title = "Periods"
    c4.y_axis.title = "Principal ($)"
    c4.style = 13
    c4.width = 18
    c4.height = 10
    yvals4 = Reference(ws, min_col=5, min_row=header_row + 1, max_row=last_row)
    c4.add_data(yvals4)
    c4.set_categories(xvals)
    c4.series[0].title = SeriesLabel(v="Principal")
    ws.add_chart(c4, f"{get_column_letter(chart_col)}{chart_row + 48}")

    # Chart 5: EndBal scatter
    c5 = ScatterChart()
    c5.title = f"{title_prefix}: End Balance"
    c5.x_axis.title = "Periods"
    c5.y_axis.title = "End Balance ($)"
    c5.style = 13
    c5.width = 18
    c5.height = 10
    yvals5 = Reference(ws, min_col=6, min_row=header_row + 1, max_row=last_row)
    c5.add_data(yvals5)
    c5.set_categories(xvals)
    c5.series[0].title = SeriesLabel(v="EndBal")
    ws.add_chart(c5, f"{get_column_letter(chart_col + 9)}{chart_row}")

    # Chart 6: Stacked bar (Interest vs Principal)
    c6 = BarChart()
    c6.type = "col"
    c6.grouping = "stacked"
    c6.title = f"{title_prefix}: Interest vs Principal"
    c6.x_axis.title = "Periods"
    c6.y_axis.title = "Amount ($)"
    c6.style = 13
    c6.width = 18
    c6.height = 10
    int_data = Reference(ws, min_col=4, min_row=header_row, max_row=last_row)
    prin_data = Reference(ws, min_col=5, min_row=header_row, max_row=last_row)
    cats = Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_row)
    c6.add_data(int_data, titles_from_data=True)
    c6.add_data(prin_data, titles_from_data=True)
    c6.set_categories(cats)
    c6.series[0].title = SeriesLabel(v="Interest Component")
    c6.series[1].title = SeriesLabel(v="Principal Repaid")
    ws.add_chart(c6, f"{get_column_letter(chart_col + 9)}{chart_row + 16}")

    # Chart 7: Proportion line chart
    c7 = LineChart()
    c7.title = f"{title_prefix}: Interest & Principal Proportion"
    c7.x_axis.title = "Periods"
    c7.y_axis.title = "Proportion"
    c7.style = 13
    c7.width = 18
    c7.height = 10
    int_pct = Reference(ws, min_col=7, min_row=header_row, max_row=last_row)
    prin_pct = Reference(ws, min_col=8, min_row=header_row, max_row=last_row)
    c7.add_data(int_pct, titles_from_data=True)
    c7.add_data(prin_pct, titles_from_data=True)
    c7.set_categories(cats)
    c7.series[0].title = SeriesLabel(v="Interest %")
    c7.series[1].title = SeriesLabel(v="Principal %")
    ws.add_chart(c7, f"{get_column_letter(chart_col + 9)}{chart_row + 32}")


def create_loan_workbook():
    """Create LoanAmortization.xlsx with all 3 test cases and 14 charts each."""
    wb = Workbook()

    for idx, tc in enumerate(TEST_CASES):
        if idx == 0:
            ws = wb.active
            ws.title = f"{tc['name']}_Constant"
        else:
            ws = wb.create_sheet(f"{tc['name']}_Constant")

        # Constant payment schedule
        df_const = constant_payment_schedule(
            tc['principal'], tc['rate'], tc['years'], tc['freq'], tc['ptype']
        )
        stats = summary_statistics(df_const)
        params = f"Principal: ${tc['principal']:,} | Rate: {tc['rate']}% | " \
                 f"Term: {tc['years']}yr | {tc['freq']} | {tc['ptype']}"

        header_row, last_row = write_schedule_to_sheet(
            ws, df_const, 1, f"Constant Payment Amortization — {tc['label']}", params
        )

        # Summary stats
        sr = last_row + 2
        ws.cell(row=sr, column=1, value="Summary Statistics").font = Font(bold=True, size=12)
        ws.cell(row=sr + 1, column=1, value="Total Periods:")
        ws.cell(row=sr + 1, column=2, value=stats['total_periods'])
        ws.cell(row=sr + 2, column=1, value="Total Interest:")
        ws.cell(row=sr + 2, column=2, value=stats['total_interest']).number_format = '#,##0.00'
        ws.cell(row=sr + 3, column=1, value="Total Payments:")
        ws.cell(row=sr + 3, column=2, value=stats['total_payments']).number_format = '#,##0.00'

        add_7_charts(ws, header_row, last_row, "Constant Payment")

        # Straight-line sheet
        ws_sl = wb.create_sheet(f"{tc['name']}_StraightLine")
        df_sl = straight_line_schedule(
            tc['principal'], tc['rate'], tc['years'], tc['freq'], tc['ptype']
        )
        stats_sl = summary_statistics(df_sl)

        header_row_sl, last_row_sl = write_schedule_to_sheet(
            ws_sl, df_sl, 1, f"Straight-Line Amortization — {tc['label']}", params
        )

        sr_sl = last_row_sl + 2
        ws_sl.cell(row=sr_sl, column=1, value="Summary Statistics").font = Font(bold=True, size=12)
        ws_sl.cell(row=sr_sl + 1, column=1, value="Total Periods:")
        ws_sl.cell(row=sr_sl + 1, column=2, value=stats_sl['total_periods'])
        ws_sl.cell(row=sr_sl + 2, column=1, value="Total Interest:")
        ws_sl.cell(row=sr_sl + 2, column=2, value=stats_sl['total_interest']).number_format = '#,##0.00'
        ws_sl.cell(row=sr_sl + 3, column=1, value="Total Payments:")
        ws_sl.cell(row=sr_sl + 3, column=2, value=stats_sl['total_payments']).number_format = '#,##0.00'

        add_7_charts(ws_sl, header_row_sl, last_row_sl, "Straight-Line")

        print(f"  {tc['label']} — {len(df_const)} periods (const), {len(df_sl)} periods (SL)")

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "after", "LoanAmortization.xlsx")
    wb.save(out)
    print(f"  Saved: {out} ({os.path.getsize(out) / 1024:.0f} KB)")
    return out


def create_portfolio_workbook():
    """Create PortfolioAnalysis.xlsx with efficient frontier and all metrics."""
    wb = Workbook()

    # Generate data
    np.random.seed(42)
    n = 61
    a1 = 100 * np.exp(np.cumsum(np.random.normal(0.008, 0.05, n)))
    a2 = 100 * np.exp(np.cumsum(np.random.normal(0.005, 0.07, n)))
    rf_prices = 100 * np.exp(np.cumsum(np.random.normal(0.002, 0.005, n)))

    analysis = full_portfolio_analysis(a1, a2, rf_prices)

    # ========================
    # Sheet 1: Efficient Frontier
    # ========================
    ws = wb.active
    ws.title = "EfficientFrontier"

    # Header
    ws.cell(row=1, column=1, value="Efficient Frontier — 44 Points").font = TITLE_FONT
    ws.merge_cells('A1:F1')

    headers = ["w(Asset1)", "w(Asset2)", "Return", "StdDev", "Variance", "Sharpe"]
    for j, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=j, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER

    frontier = analysis['frontier']
    for i, row in frontier.iterrows():
        r = 4 + i
        ws.cell(row=r, column=1, value=row['w_asset1']).number_format = '0.0000'
        ws.cell(row=r, column=2, value=row['w_asset2']).number_format = '0.0000'
        ws.cell(row=r, column=3, value=row['Return']).number_format = '0.000000'
        ws.cell(row=r, column=4, value=row['StdDev']).number_format = '0.000000'
        ws.cell(row=r, column=5, value=row['Variance']).number_format = '0.000000'
        sharpe = (row['Return'] - analysis['rf_rate']) / row['StdDev'] if row['StdDev'] > 0 else 0
        ws.cell(row=r, column=6, value=sharpe).number_format = '0.0000'
        for j in range(1, 7):
            ws.cell(row=r, column=j).border = THIN_BORDER

    last_row = 3 + len(frontier)

    # Efficient Frontier chart
    c1 = ScatterChart()
    c1.title = "Efficient Frontier"
    c1.x_axis.title = "Standard Deviation"
    c1.y_axis.title = "Expected Return"
    c1.style = 13
    c1.width = 22
    c1.height = 14
    xvals = Reference(ws, min_col=4, min_row=4, max_row=last_row)
    yvals = Reference(ws, min_col=3, min_row=4, max_row=last_row)
    c1.add_data(yvals)
    c1.set_categories(xvals)
    c1.series[0].title = SeriesLabel(v="Frontier")
    c1.series[0].graphicalProperties.line.noFill = True
    ws.add_chart(c1, "H3")

    # Sharpe ratio chart
    c2 = LineChart()
    c2.title = "Sharpe Ratio vs Weight"
    c2.x_axis.title = "w(Asset 1)"
    c2.y_axis.title = "Sharpe Ratio"
    c2.style = 13
    c2.width = 22
    c2.height = 14
    xvals2 = Reference(ws, min_col=1, min_row=4, max_row=last_row)
    yvals2 = Reference(ws, min_col=6, min_row=3, max_row=last_row)
    c2.add_data(yvals2, titles_from_data=True)
    c2.set_categories(xvals2)
    ws.add_chart(c2, "H22")

    widths = [12, 12, 14, 14, 14, 12]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    # ========================
    # Sheet 2: Key Metrics
    # ========================
    ws2 = wb.create_sheet("Metrics")
    ws2.cell(row=1, column=1, value="Portfolio Analysis — Key Metrics").font = TITLE_FONT
    ws2.merge_cells('A1:D1')

    metrics = [
        ("Asset 1 Expected Return", analysis['expected_returns'][0], '0.0000%'),
        ("Asset 2 Expected Return", analysis['expected_returns'][1], '0.0000%'),
        ("Risk-Free Rate", analysis['rf_rate'], '0.0000%'),
        ("", "", ""),
        ("MVP Weight (Asset 1)", analysis['mvp_weights'][0], '0.0000'),
        ("MVP Weight (Asset 2)", analysis['mvp_weights'][1], '0.0000'),
        ("MVP Return", analysis['mvp_return'], '0.0000%'),
        ("MVP Std Dev", analysis['mvp_std'], '0.0000%'),
        ("", "", ""),
        ("ORP Weight (Asset 1)", analysis['orp_weights'][0], '0.0000'),
        ("ORP Weight (Asset 2)", analysis['orp_weights'][1], '0.0000'),
        ("ORP Return", analysis['orp_return'], '0.0000%'),
        ("ORP Std Dev", analysis['orp_std'], '0.0000%'),
        ("ORP Sharpe", analysis['orp_sharpe'], '0.0000'),
    ]

    for i, (label, val, fmt) in enumerate(metrics, 3):
        ws2.cell(row=i, column=1, value=label).font = Font(bold=True, size=11)
        if val != "":
            cell = ws2.cell(row=i, column=2, value=val)
            cell.number_format = fmt
            cell.font = Font(size=11)

    ws2.column_dimensions['A'].width = 28
    ws2.column_dimensions['B'].width = 16

    # VCV Matrix
    ws2.cell(row=20, column=1, value="Variance-Covariance Matrix (Annualized)").font = Font(
        bold=True, size=12)
    vcv = analysis['vcv_matrix']
    ws2.cell(row=21, column=2, value="Asset 1").font = HEADER_FONT
    ws2.cell(row=21, column=2).fill = HEADER_FILL
    ws2.cell(row=21, column=3, value="Asset 2").font = HEADER_FONT
    ws2.cell(row=21, column=3).fill = HEADER_FILL
    ws2.cell(row=22, column=1, value="Asset 1").font = Font(bold=True)
    ws2.cell(row=22, column=2, value=vcv[0, 0]).number_format = '0.000000'
    ws2.cell(row=22, column=3, value=vcv[0, 1]).number_format = '0.000000'
    ws2.cell(row=23, column=1, value="Asset 2").font = Font(bold=True)
    ws2.cell(row=23, column=2, value=vcv[1, 0]).number_format = '0.000000'
    ws2.cell(row=23, column=3, value=vcv[1, 1]).number_format = '0.000000'

    # Correlation Matrix
    corr = analysis['correlation_matrix']
    ws2.cell(row=26, column=1, value="Correlation Matrix").font = Font(bold=True, size=12)
    ws2.cell(row=27, column=2, value="Asset 1").font = HEADER_FONT
    ws2.cell(row=27, column=2).fill = HEADER_FILL
    ws2.cell(row=27, column=3, value="Asset 2").font = HEADER_FONT
    ws2.cell(row=27, column=3).fill = HEADER_FILL
    ws2.cell(row=28, column=1, value="Asset 1").font = Font(bold=True)
    ws2.cell(row=28, column=2, value=corr[0, 0]).number_format = '0.000000'
    ws2.cell(row=28, column=3, value=corr[0, 1]).number_format = '0.000000'
    ws2.cell(row=29, column=1, value="Asset 2").font = Font(bold=True)
    ws2.cell(row=29, column=2, value=corr[1, 0]).number_format = '0.000000'
    ws2.cell(row=29, column=3, value=corr[1, 1]).number_format = '0.000000'

    # Weights comparison bar chart
    ws2.cell(row=32, column=1, value="Portfolio").font = Font(bold=True)
    ws2.cell(row=32, column=2, value="Asset 1").font = HEADER_FONT
    ws2.cell(row=32, column=2).fill = HEADER_FILL
    ws2.cell(row=32, column=3, value="Asset 2").font = HEADER_FONT
    ws2.cell(row=32, column=3).fill = HEADER_FILL
    ws2.cell(row=33, column=1, value="MVP")
    ws2.cell(row=33, column=2, value=analysis['mvp_weights'][0]).number_format = '0.0000'
    ws2.cell(row=33, column=3, value=analysis['mvp_weights'][1]).number_format = '0.0000'
    ws2.cell(row=34, column=1, value="ORP")
    ws2.cell(row=34, column=2, value=analysis['orp_weights'][0]).number_format = '0.0000'
    ws2.cell(row=34, column=3, value=analysis['orp_weights'][1]).number_format = '0.0000'

    c3 = BarChart()
    c3.type = "col"
    c3.grouping = "clustered"
    c3.title = "Portfolio Weights: MVP vs ORP"
    c3.style = 13
    c3.width = 16
    c3.height = 10
    data = Reference(ws2, min_col=2, min_row=32, max_col=3, max_row=34)
    cats = Reference(ws2, min_col=1, min_row=33, max_row=34)
    c3.add_data(data, titles_from_data=True)
    c3.set_categories(cats)
    ws.add_chart(c3, "H40")

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "after", "PortfolioAnalysis.xlsx")
    wb.save(out)
    print(f"  Saved: {out} ({os.path.getsize(out) / 1024:.0f} KB)")
    return out


def main():
    print("=" * 60)
    print("GENERATING EXCEL WORKBOOKS")
    print("=" * 60)

    print("\n1. Loan Amortization Workbook")
    loan_path = create_loan_workbook()

    print("\n2. Portfolio Analysis Workbook")
    port_path = create_portfolio_workbook()

    print(f"\n{'=' * 60}")
    print("Excel workbooks generated successfully!")
    print(f"  Loan: {loan_path}")
    print(f"  Portfolio: {port_path}")

    return loan_path, port_path


if __name__ == "__main__":
    main()
